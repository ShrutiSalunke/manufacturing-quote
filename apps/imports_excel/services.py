"""Excel template generation and import services."""
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.contrib.contenttypes.models import ContentType
from django.db import transaction
from openpyxl import Workbook, load_workbook

from apps.catalog.models import (
    CustomFieldDefinition,
    CustomFieldValue,
    LaborRole,
    Machine,
    Material,
)
from apps.core.logging_utils import log_event
from apps.core.models import Plant
from apps.imports_excel.models import ImportJob, ImportRowError
from apps.quotes.models import Quote, QuoteLine, QuoteLineParameterValue


BASE_COLUMNS = {
    ImportJob.EntityType.MATERIAL: [
        "plant_code",
        "code",
        "name",
        "uom",
        "density",
        "category",
        "unit_price",
        "currency",
        "effective_from",
        "effective_to",
        "is_active",
    ],
    ImportJob.EntityType.MACHINE: [
        "plant_code",
        "code",
        "name",
        "hourly_rate",
        "setup_rate",
        "efficiency_percent",
        "is_active",
    ],
    ImportJob.EntityType.LABOR: [
        "plant_code",
        "code",
        "name",
        "hourly_rate",
        "is_active",
    ],
    ImportJob.EntityType.QUOTE_INPUT: [
        "quote_number",
        "line_no",
        "parameter_code",
        "value",
    ],
}

ENTITY_TO_CUSTOM = {
    ImportJob.EntityType.MATERIAL: CustomFieldDefinition.EntityType.MATERIAL,
    ImportJob.EntityType.MACHINE: CustomFieldDefinition.EntityType.MACHINE,
    ImportJob.EntityType.LABOR: CustomFieldDefinition.EntityType.LABOR,
    ImportJob.EntityType.QUOTE_INPUT: CustomFieldDefinition.EntityType.QUOTE_INPUT,
}

EXAMPLE_ROWS = {
    ImportJob.EntityType.MATERIAL: {
        "plant_code": "MAIN",
        "code": "MS-SHEET",
        "name": "Mild Steel Sheet",
        "uom": "kg",
        "density": "7.85",
        "category": "Metal",
        "unit_price": "65",
        "currency": "INR",
        "effective_from": "2024-01-01",
        "effective_to": "",
        "is_active": "TRUE",
    },
    ImportJob.EntityType.MACHINE: {
        "plant_code": "MAIN",
        "code": "LASER-1",
        "name": "Fiber Laser Cutter",
        "hourly_rate": "1200",
        "setup_rate": "500",
        "efficiency_percent": "90",
        "is_active": "TRUE",
    },
    ImportJob.EntityType.LABOR: {
        "plant_code": "MAIN",
        "code": "OPERATOR",
        "name": "Machine Operator",
        "hourly_rate": "250",
        "is_active": "TRUE",
    },
    ImportJob.EntityType.QUOTE_INPUT: {
        "quote_number": "Q-EXAMPLE",
        "line_no": "1",
        "parameter_code": "LENGTH",
        "value": "500",
    },
}


def get_importable_custom_fields(entity_type):
    custom_type = ENTITY_TO_CUSTOM.get(entity_type)
    if not custom_type:
        return []
    return list(
        CustomFieldDefinition.objects.filter(
            entity_type=custom_type, is_active=True, is_importable=True
        ).order_by("sort_order", "key")
    )


def build_headers(entity_type):
    headers = list(BASE_COLUMNS[entity_type])
    custom = get_importable_custom_fields(entity_type)
    for cf in custom:
        headers.append(cf.import_column_header or cf.key)
    return headers, custom


def generate_template_workbook(entity_type) -> bytes:
    headers, custom = build_headers(entity_type)
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(headers)
    example = dict(EXAMPLE_ROWS.get(entity_type, {}))
    for cf in custom:
        example[cf.import_column_header or cf.key] = ""
    ws.append([example.get(h, "") for h in headers])

    instr = wb.create_sheet("INSTRUCTIONS")
    instr.append(["Column", "Required", "Notes"])
    for h in BASE_COLUMNS[entity_type]:
        instr.append([h, "Yes", "System column"])
    for cf in custom:
        instr.append(
            [
                cf.import_column_header or cf.key,
                "Yes" if cf.is_required else "No",
                f"Custom field ({cf.data_type}) — key={cf.key}",
            ]
        )
    instr.append([])
    instr.append(["Upsert key for masters:", "plant_code + code"])
    instr.append(["Boolean values:", "TRUE / FALSE / 1 / 0"])
    instr.append(["Dates:", "YYYY-MM-DD"])
    instr.append(["Units:", "Use plant currency; rates are per hour; density typically g/cm3 or kg/L"])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse_bool(val, default=True):
    if val is None or str(val).strip() == "":
        return default
    return str(val).strip().lower() in ("1", "true", "yes", "y")


def _parse_date(val):
    if val is None or str(val).strip() == "":
        return None
    if isinstance(val, datetime):
        return val.date()
    return datetime.strptime(str(val).strip()[:10], "%Y-%m-%d").date()


def _parse_decimal(val, field_name):
    if val is None or str(val).strip() == "":
        return None
    try:
        return Decimal(str(val))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"Invalid number for {field_name}: {val}") from exc


def _row_dict(headers, row):
    data = {}
    for i, h in enumerate(headers):
        data[h] = row[i] if i < len(row) else None
    return data


def _set_custom_values(obj, custom_fields, row_data):
    ct = ContentType.objects.get_for_model(obj)
    for cf in custom_fields:
        header = cf.import_column_header or cf.key
        raw = row_data.get(header)
        if (raw is None or str(raw).strip() == "") and cf.is_required:
            raise ValueError(f"Required custom field missing: {header}")
        if raw is None or str(raw).strip() == "":
            continue
        cfv, _ = CustomFieldValue.objects.get_or_create(
            definition=cf, content_type=ct, object_id=obj.pk
        )
        if cf.data_type == CustomFieldDefinition.DataType.NUMBER:
            cfv.set_value(_parse_decimal(raw, header))
        elif cf.data_type == CustomFieldDefinition.DataType.DATE:
            cfv.set_value(_parse_date(raw))
        elif cf.data_type == CustomFieldDefinition.DataType.BOOL:
            cfv.set_value(_parse_bool(raw, False))
        else:
            cfv.set_value(str(raw))
        cfv.save()


def import_workbook(job: ImportJob, *, request=None):
    headers_expected, custom_fields = build_headers(job.entity_type)
    success = 0
    failed = 0
    errors = []

    try:
        wb = load_workbook(job.uploaded_file, data_only=True)
        ws = wb["Data"] if "Data" in wb.sheetnames else wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise ValueError("Empty workbook")
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        for h in BASE_COLUMNS[job.entity_type]:
            if h not in headers:
                raise ValueError(f"Missing required column: {h}")

        for idx, row in enumerate(rows[1:], start=2):
            if row is None or all(c is None or str(c).strip() == "" for c in row):
                continue
            data = _row_dict(headers, row)
            try:
                with transaction.atomic():
                    _import_one_row(job.entity_type, data, custom_fields)
                success += 1
            except Exception as row_exc:
                failed += 1
                errors.append(
                    ImportRowError(
                        job=job,
                        row_number=idx,
                        column="",
                        message=str(row_exc),
                    )
                )

        job.row_success = success
        job.row_failed = failed
        if failed and success:
            job.status = ImportJob.Status.PARTIAL
        elif failed and not success:
            job.status = ImportJob.Status.FAILED
        else:
            job.status = ImportJob.Status.SUCCESS
        job.error_summary = f"{success} succeeded, {failed} failed"
        job.save()
        if errors:
            ImportRowError.objects.bulk_create(errors)

        if failed:
            log_event(
                "ERROR" if not success else "WARNING",
                "imports_excel",
                "IMPORT",
                f"Import {job.entity_type} finished with failures: {job.error_summary}",
                request=request,
                user=job.created_by,
                context={
                    "import_job_id": job.pk,
                    "entity_type": job.entity_type,
                    "row_failed": failed,
                    "row_success": success,
                    "errors": [{"row": e.row_number, "message": e.message} for e in errors[:50]],
                },
                correlation_id=job.correlation_id,
            )
        return job
    except Exception as exc:
        job.status = ImportJob.Status.FAILED
        job.error_summary = str(exc)
        job.row_success = success
        job.row_failed = failed
        job.save()
        entry = log_event(
            "ERROR",
            "imports_excel",
            "IMPORT",
            f"Import job failed: {exc}",
            request=request,
            user=job.created_by,
            exc=exc,
            context={"import_job_id": job.pk, "entity_type": job.entity_type},
            correlation_id=job.correlation_id,
        )
        if entry:
            job.correlation_id = entry.correlation_id
            job.save(update_fields=["correlation_id"])
        raise


def _import_one_row(entity_type, data, custom_fields):
    if entity_type == ImportJob.EntityType.MATERIAL:
        plant = Plant.objects.get(code=str(data["plant_code"]).strip())
        code = str(data["code"]).strip()
        obj, _ = Material.objects.update_or_create(
            plant=plant,
            code=code,
            defaults={
                "name": str(data.get("name") or code),
                "uom": str(data.get("uom") or "kg"),
                "density": _parse_decimal(data.get("density"), "density"),
                "category": str(data.get("category") or ""),
                "unit_price": _parse_decimal(data.get("unit_price"), "unit_price") or Decimal("0"),
                "currency": str(data.get("currency") or plant.currency),
                "effective_from": _parse_date(data.get("effective_from")),
                "effective_to": _parse_date(data.get("effective_to")),
                "is_active": _parse_bool(data.get("is_active"), True),
            },
        )
        _set_custom_values(obj, custom_fields, data)
        return obj

    if entity_type == ImportJob.EntityType.MACHINE:
        plant = Plant.objects.get(code=str(data["plant_code"]).strip())
        code = str(data["code"]).strip()
        obj, _ = Machine.objects.update_or_create(
            plant=plant,
            code=code,
            defaults={
                "name": str(data.get("name") or code),
                "hourly_rate": _parse_decimal(data.get("hourly_rate"), "hourly_rate") or Decimal("0"),
                "setup_rate": _parse_decimal(data.get("setup_rate"), "setup_rate") or Decimal("0"),
                "efficiency_percent": _parse_decimal(data.get("efficiency_percent"), "efficiency_percent")
                or Decimal("100"),
                "is_active": _parse_bool(data.get("is_active"), True),
            },
        )
        _set_custom_values(obj, custom_fields, data)
        return obj

    if entity_type == ImportJob.EntityType.LABOR:
        plant = Plant.objects.get(code=str(data["plant_code"]).strip())
        code = str(data["code"]).strip()
        obj, _ = LaborRole.objects.update_or_create(
            plant=plant,
            code=code,
            defaults={
                "name": str(data.get("name") or code),
                "hourly_rate": _parse_decimal(data.get("hourly_rate"), "hourly_rate") or Decimal("0"),
                "is_active": _parse_bool(data.get("is_active"), True),
            },
        )
        _set_custom_values(obj, custom_fields, data)
        return obj

    if entity_type == ImportJob.EntityType.QUOTE_INPUT:
        quote = Quote.objects.get(number=str(data["quote_number"]).strip())
        if not quote.is_editable:
            raise ValueError("Quote is issued and locked")
        line_no = int(data["line_no"])
        lines = list(quote.lines.order_by("sort_order", "id"))
        if line_no < 1 or line_no > len(lines):
            raise ValueError(f"Invalid line_no {line_no}")
        line = lines[line_no - 1]
        code = str(data["parameter_code"]).strip().upper()
        QuoteLineParameterValue.objects.update_or_create(
            line=line,
            parameter_code=code,
            defaults={"value": str(data.get("value") if data.get("value") is not None else "")},
        )
        return line

    raise ValueError(f"Unknown entity type {entity_type}")
